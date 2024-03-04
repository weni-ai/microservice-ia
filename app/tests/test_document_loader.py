import unittest
from app.loaders.loaders import (
    DataLoader,
    DataLoaderCls,
    PDFLoader,
    DocxLoader,
    TxtLoader,
    URLsLoader,
    XlsxLoader,
    pdf_loader,
    txt_loader,
    docx_loader,
    u_docx_loader,
    xlsx_loader
)
from unittest import mock
from app.loaders import (
    load_file_url_and_split_text,
    load_file_url_and_get_pages_text,
    load_file_and_get_raw_text,
    load_file_url_and_get_raw_text
)
import os
import pathlib
import shutil
import xlsxwriter
from docx import Document
from reportlab.pdfgen.canvas import Canvas
from app.text_splitters import TextSplitter, character_text_splitter
from typing import List


TEST_DIR = pathlib.Path(__file__).parent.resolve()
TEST_FILE_PATH = os.path.join(TEST_DIR, 'test_files')


class TestDocumentLoader(unittest.TestCase):
    test_dir = TEST_DIR
    path = TEST_FILE_PATH
    text_string = 'Hello, World!'
    file_name = 'test_file'
    text_splitter = TextSplitter(character_text_splitter())

    @staticmethod
    def _create_test_dir():
        if not os.path.exists(TestDocumentLoader.path):
            os.mkdir(TestDocumentLoader.path)

    @staticmethod
    def _create_pdf(path: str, name: str, text: str):
        file_path = f'{path}/{name}.pdf'
        canvas = Canvas(file_path)
        canvas.drawString(72, 72, text)
        canvas.save()

    @staticmethod
    def _create_txt(path: str, name: str, text: str):
        file_path = f'{path}/{name}.txt'
        with open(file_path, 'w') as f:
            f.write(text)

    @staticmethod
    def _create_docx(path: str, name: str, text: str):
        file_path = f'{path}/{name}.docx'
        document = Document()
        document.add_paragraph(text, style='Intense Quote')
        document.save(file_path)

    @staticmethod
    def _create_xlsx(path: str, name: str):
        file_path = f'{path}/{name}.xlsx'
        workbook = xlsxwriter.Workbook(file_path)
        worksheet = workbook.add_worksheet("Test sheet")
        scores = (
            ['Lorem', 4576],
            ['Ipsum',   345],
            ['Dolor',  9088],
            ['Sit',    88],
            ['Amet',    15],
        )
        row = 0
        col = 0
        for name, score in (scores):
            worksheet.write(row, col, name)
            worksheet.write(row, col + 1, score)
            row += 1
        workbook.close()

    @classmethod
    def setUpClass(cls) -> None:
        cls._create_test_dir()
        cls._create_pdf(cls.path, cls.file_name, cls.text_string)
        cls._create_txt(cls.path, cls.file_name, cls.text_string)
        cls._create_docx(cls.path, cls.file_name, cls.text_string)
        cls._create_xlsx(cls.path, cls.file_name)

    def test_load_pdf(self):
        file_path = f'{self.path}/{self.file_name}.pdf'
        data_loader = DataLoader(pdf_loader, file_path)
        raw_text = data_loader.raw_text()
        self.assertEqual(raw_text, self.text_string.lower())

    def test_load_txt(self):
        file_path = f'{self.path}/{self.file_name}.txt'
        data_loader = DataLoader(txt_loader, file_path)
        raw_text = data_loader.raw_text()
        self.assertEqual(raw_text, self.text_string.lower())

    def test_load_udocx(self):
        file_path = f'{self.path}/{self.file_name}.docx'
        data_loader = DataLoader(u_docx_loader, file_path)
        raw_text = data_loader.raw_text()
        self.assertEqual(raw_text, self.text_string.lower())

    def test_load_docx(self):
        file_path = f'{self.path}/{self.file_name}.docx'
        data_loader = DataLoader(docx_loader, file_path)
        raw_text = data_loader.raw_text()
        self.assertEqual(raw_text, self.text_string.lower())

    def test_load_xlsx(self):
        file_path = f'{self.path}/{self.file_name}.xlsx'
        data_loader = DataLoader(xlsx_loader, file_path)
        raw_text = data_loader.raw_text()
        self.assertEqual(type(raw_text), str)

    @mock.patch("app.loaders.loaders.XlsxLoader._get_temp_file")
    def test_load_xlsx_cls(self, mock_file_url):
        file_path = f'{self.path}/{self.file_name}.xlsx'
        mock_file_url.return_value = (file_path, "")
        xlsx_loader = XlsxLoader(file_path)
        split_pages: List[Document] = xlsx_loader.load_and_split_text(self.text_splitter)
        self.assertEqual(list, type(split_pages))

    def test_pdf_loader_cls(self):
        file_path = f'{self.path}/{self.file_name}.pdf'
        pdf_loader = PDFLoader(file_path)
        split_pages: List[Document] = pdf_loader.load_and_split_text(self.text_splitter)
        self.assertEqual(list, type(split_pages))
    
    def test_urls_loader_cls(self):
        urls_loader = URLsLoader("https://en.wikipedia.org/wiki/Unit_testing")
        split_pages: List[Document] = urls_loader.load()
        self.assertEqual(list, type(split_pages))

    def test_urls_loader_and_split_cls(self):
        urls_loader = URLsLoader("https://en.wikipedia.org/wiki/Unit_testing")
        split_pages: List[Document] = urls_loader.load_and_split_text(self.text_splitter)
        self.assertEqual(list, type(split_pages))
    
    def test_urls_list_loader_and_split_cls(self):
        urls = ["https://en.wikipedia.org/wiki/Unit_testing"]
        urls_loader = URLsLoader(urls)
        split_pages: List[Document] = urls_loader.load_and_split_text(self.text_splitter)
        self.assertEqual(list, type(split_pages))

    def test_docx_loader_cls(self):
        file_path = f'{self.path}/{self.file_name}.docx'
        docx_loader = DocxLoader(file_path)
        split_pages: List[Document] = docx_loader.load_and_split_text(self.text_splitter)
        self.assertEqual(list, type(split_pages))
    
    def test_txt_loader_cls(self):
        file_path = f'{self.path}/{self.file_name}.txt'
        docx_loader = TxtLoader(file_path)
        split_pages: List[Document] = docx_loader.load_and_split_text(self.text_splitter)
        self.assertEqual(list, type(split_pages))

    def test_load_file_url_and_split_text(self):
        file_path = f'{self.path}/{self.file_name}.pdf'
        file_type = "pdf"
        docs = load_file_url_and_split_text(file_path, file_type, self.text_splitter)
        self.assertEqual(list, type(docs))

    def test_load_file_url_and_get_pages_text(self):  # this function is deprecated
        file_path = f'{self.path}/{self.file_name}.pdf'
        file_type = "pdf"
        docs = load_file_url_and_get_pages_text(file_path ,file_type)
        self.assertEqual(list, type(docs))

    @mock.patch.dict(os.environ, {"FILE_PATH": TEST_FILE_PATH})
    def test_load_file_and_get_raw_text(self):  # this function is deprecated
        file_type = "pdf"
        raw_text = load_file_and_get_raw_text(f"{self.file_name}.{file_type}" ,file_type)
        self.assertEqual(str, type(raw_text))

    def test_file_url_and_get_raw_text(self):  # this function is deprecated
        file_path = f'{self.path}/{self.file_name}.pdf'
        file_type = "pdf"
        raw_text = load_file_url_and_get_raw_text(file_path ,file_type)
        self.assertEqual(str, type(raw_text))

    def test_pdf_loader_cls_raw_text(self):
        file_path = f'{self.path}/{self.file_name}.pdf'
        pdf_loader = PDFLoader(file_path)
        raw_text = pdf_loader.raw_text()
        self.assertEqual(str, type(raw_text))

    def test_data_loader_cls(self):
        file_path = f'{self.path}/{self.file_name}.pdf'
        data_loader = DataLoaderCls(PDFLoader, file_path)
        loaded_data: List[Document] = data_loader.load()
        self.assertEqual(list, type(loaded_data))
    
    def test_data_loader_raw_text(self):
        file_path = f'{self.path}/{self.file_name}.pdf'
        data_loader = DataLoaderCls(PDFLoader, file_path)
        loaded_data: str = data_loader.raw_text()
        self.assertEqual(str, type(loaded_data))


    @classmethod
    def tearDownClass(cls):
        if os.path.exists(TestDocumentLoader.path):
            shutil.rmtree(TestDocumentLoader.path, ignore_errors=True)
